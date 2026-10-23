"""Session-only spreadsheet loading, explicit schema mapping, and append validation."""
from dataclasses import dataclass
from difflib import get_close_matches
from io import BytesIO, StringIO
import json
import csv
import time
import pandas as pd
from openpyxl import load_workbook

MAX_FILES = 10


def validate_limits(file_count: int, record_count: int = 0) -> None:
    """Enforce the file-count limit; record counts are currently unrestricted."""
    if file_count > MAX_FILES:
        raise ValueError('Upload at most 10 files. Remove files before continuing.')


def preview_rows(source: 'Source', count: int = 10) -> pd.DataFrame:
    """Return the first records without changing source data."""
    if count < 1:
        raise ValueError('Preview row count must be positive.')
    return source.frame.head(count)


@dataclass
class Source:
    name: str
    sheet: str
    header: int
    frame: pd.DataFrame
    header_end: int | None = None

    @property
    def key(self) -> str:
        """Identify a dataset without conflating worksheets from one workbook."""
        return json.dumps([self.name, self.sheet]) if self.sheet else self.name


def worksheets(content):
    return pd.ExcelFile(BytesIO(content), engine="openpyxl").sheet_names


def load(content, name, sheet=None, header=None):
    if name.lower().endswith('.csv'):
        rows = list(csv.reader(StringIO(content.decode('utf-8-sig'))))
        raw = pd.DataFrame(rows, dtype=object).replace('', pd.NA)
    else:
        raw = pd.read_excel(BytesIO(content), sheet_name=sheet or 0, header=None, dtype=object)
    populated = raw.dropna(how='all')
    if populated.empty:
        raise ValueError('The file has no populated rows.')
    positions = [int(populated.index[0])] if header is None else ([header] if isinstance(header, int) and not isinstance(header, bool) else header)
    if (not isinstance(positions, list) or not positions or
            any(type(p) is not int for p in positions) or
            positions != list(range(positions[0], positions[-1] + 1))):
        raise ValueError('Header rows must be a nonempty consecutive range of row positions.')
    position, end = positions[0], positions[-1]
    if position < 0 or end >= len(raw):
        raise ValueError('Header row is outside the worksheet.')
    header_cells = raw.iloc[position:end + 1].copy()
    if len(positions) > 1 and name.lower().endswith('.xlsx'):
        # Expand only actual Excel merges, never guess groups by filling arbitrary blanks.
        book = load_workbook(BytesIO(content), data_only=True)
        try:
            ws = book[sheet] if sheet else book.worksheets[0]
            for merged in ws.merged_cells.ranges:
                if merged.min_row >= position + 1 and merged.max_row <= end + 1:
                    value = ws.cell(merged.min_row, merged.min_col).value
                    for row in range(merged.min_row, merged.max_row + 1):
                        for col in range(merged.min_col, min(merged.max_col, len(raw.columns)) + 1):
                            header_cells.iat[row - position - 1, col - 1] = value
        finally:
            book.close()
    names = []
    for col in header_cells:
        parts = []
        for value in header_cells[col]:
            part = '' if pd.isna(value) else ' '.join(str(value).split())
            if part and (not parts or part != parts[-1]):
                parts.append(part)
        names.append(' '.join(parts))
    if any(not n for n in names):
        raise ValueError('Blank header detected. Choose a different header row or repair the file.')
    if len(set(names)) != len(names):
        raise ValueError('Duplicate headers detected after merging and normalizing. Adjust the header rows or repair the source headers.')
    frame = raw.iloc[end + 1:].copy().dropna(how='all')
    frame.columns = names
    frame.index = frame.index + 1  # actual one-based worksheet row numbers
    return Source(name, sheet or '', position + 1, frame, end + 1)


def default_config():
    return {'version': 1, 'name': '', 'columns': [], 'approved': {},
            'rejected': {}, 'mappings': {}, 'missing': 'Insert blanks', 'trace': False}


def validate_config(config):
    if config.get('export_format', 'xlsx') not in ['csv', 'xlsx']:
        raise ValueError('Output file type must be csv or xlsx.')
    if config.get('version') != 1:
        raise ValueError('Unsupported configuration version.')
    cols = config.get('columns', [])
    names = [c['name'] for c in cols]
    if not names or any(not isinstance(n, str) or not n.strip() for n in names) or len(set(names)) != len(names):
        raise ValueError('Master column names must be nonempty and unique.')
    if any(c.get('type') not in ['Text', 'Date', 'Number'] or not isinstance(c.get('required'), bool) for c in cols):
        raise ValueError('Invalid column type or required flag.')
    if set(names) & {'Source File', 'Source Worksheet', 'Source Row'}:
        raise ValueError('Source File, Source Worksheet and Source Row are reserved names.')
    for key in ['approved', 'rejected', 'mappings']:
        if not isinstance(config.get(key), dict):
            raise ValueError(f'{key} must be a dictionary.')
    if config.get('missing') not in ['Insert blanks', 'Stop'] or not isinstance(config.get('trace'), bool):
        raise ValueError('Invalid missing-column or traceability setting.')
    return config


def read_config(content):
    return validate_config(json.loads(content))


def propose(header, config):
    names = [c['name'] for c in config['columns']]
    rejected = config['rejected'].get(header, [])
    if header in names and header not in rejected:
        return header, 'Exact match'
    approved = config['approved'].get(header)
    if approved in names and approved not in rejected:
        return approved, 'Approved synonym'
    candidates = [n for n in names if n not in rejected]
    matches = get_close_matches(header.lower().replace('_', ' '), [n.lower() for n in candidates], n=1, cutoff=.45)
    suggestion = next((n for n in candidates if matches and n.lower() == matches[0]), None)
    return None, f'Suggested: {suggestion}' if suggestion else 'Needs review'


def consolidate(sources, config, mappings):
    validate_limits(len({s.name for s in sources}), sum(len(s.frame) for s in sources))
    if not sources:
        raise ValueError('Upload at least one dataset before consolidating.')
    validate_config(config)
    started = time.perf_counter()
    issues, frames, counts = [], [], []
    names = [c['name'] for c in config['columns']]

    def issue(source, col, kind, description, severity='Warning', row=None):
        issues.append({'Source filename': source.name, 'Worksheet': source.sheet, 'Source row': row,
                       'Source column': col, 'Issue type': kind, 'Description': description,
                       'Severity': severity, 'Resolution status': 'Needs review' if severity == 'Error' else 'Reported; record retained'})

    for source in sources:
        mapping = mappings.get(source.key, {})
        if not mapping and sum(s.name == source.name for s in sources) == 1:
            mapping = mappings.get(source.name, {})  # older single-sheet templates
        targets = [v for v in mapping.values() if v and v != '__IGNORE__']
        if len(targets) != len(set(targets)):
            issue(source, '', 'Conflicting mapping', 'Multiple source fields map to one master column.', 'Error')
        for col in source.frame.columns:
            target = mapping.get(col)
            if not target:
                issue(source, col, 'Unreviewed header', 'Map this field or explicitly ignore it.', 'Error')
            elif target != '__IGNORE__' and target not in names:
                issue(source, col, 'Invalid mapping', 'Target is absent from the master schema.', 'Error')
            elif target in config['rejected'].get(col, []):
                issue(source, col, 'Rejected synonym', 'Change the rejected decision before using this mapping.', 'Error')
        out = pd.DataFrame(index=source.frame.index)
        for col in config['columns']:
            name = col['name']
            origin = next((k for k, v in mapping.items() if v == name and k in source.frame.columns), None)
            if origin is None:
                issue(source, name, 'Missing column', 'Missing values inserted.' if config['missing'] == 'Insert blanks' else 'Missing-column policy requires a stop.',
                      'Error' if col['required'] or config['missing'] == 'Stop' else 'Warning')
                out[name] = pd.NA
                continue
            values = source.frame[origin].copy()
            if col['type'] != 'Text':
                converted = pd.to_numeric(values, errors='coerce') if col['type'] == 'Number' else pd.to_datetime(values, format='ISO8601', errors='coerce')
                for row in values.index[values.notna() & converted.isna()]:
                    issue(source, origin, 'Invalid ' + col['type'].lower(), f'Original value retained: {values.loc[row]}', row=int(row))
                values = converted.astype(object).where(converted.notna(), values)
            if col['required']:
                for row in values.index[values.isna() | values.astype(str).str.strip().eq('')]:
                    issue(source, origin, 'Missing required value', 'Required value is empty.', 'Error', int(row))
            out[name] = values
        if config['trace']:
            out['Source File'], out['Source Worksheet'], out['Source Row'] = source.name, source.sheet, source.frame.index
        frames.append(out)
        counts.append({'File': source.name, 'Worksheet': source.sheet, 'Rows read': len(source.frame), 'Rows included': len(out)})
    blocked = any(i['Severity'] == 'Error' for i in issues)
    master = pd.concat(frames, ignore_index=True) if frames and not blocked else pd.DataFrame(columns=names)
    if not blocked and not master.empty:
        duplicates = master.duplicated(subset=names, keep=False)
        for index in master.index[duplicates]:
            values = master.loc[index]
            issues.append({'Source filename': values.get('Source File', ''), 'Worksheet': values.get('Source Worksheet', ''),
                           'Source row': values.get('Source Row', None), 'Source column': '', 'Issue type': 'Duplicate record',
                           'Description': 'Identical master-field values; record retained.', 'Severity': 'Warning', 'Resolution status': 'Reported; record retained'})
    if not blocked:
        assert len(master) == sum(len(s.frame) for s in sources), 'Row reconciliation failed'
    else:
        for count in counts:
            count['Rows included'] = 0
    report = pd.DataFrame(issues, columns=['Source filename', 'Worksheet', 'Source row', 'Source column', 'Issue type', 'Description', 'Severity', 'Resolution status'])
    summary = {'files': len({s.name for s in sources}), 'datasets': len(sources), 'input_rows': sum(len(s.frame) for s in sources), 'output_rows': len(master),
               'columns': list(master.columns), 'blocked': blocked, 'warnings': sum(i['Severity'] == 'Warning' for i in issues),
               'errors': sum(i['Severity'] == 'Error' for i in issues), 'duration_seconds': round(time.perf_counter() - started, 4), 'sources': counts}
    return master, report, summary


def excel_bytes(master, exceptions):
    buffer = BytesIO()
    # Disable spreadsheet formulas in untrusted string cells.
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for name, frame in [('Master', master), ('Exceptions', exceptions)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            for row in writer.sheets[name]:
                for cell in row:
                    if cell.data_type == 'f':
                        cell.data_type = 's'
    return buffer.getvalue()
