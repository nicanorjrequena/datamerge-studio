"""User-selected multi-row header normalization and attribution."""
from io import BytesIO
import json
import pandas as pd
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
from src.engine import load, default_config, read_config, consolidate


def test_csv_header_range_whitespace_and_row_preservation():
    source = load(b'Report,,\n Customer ,Sales,Sales\n Name , Amount ,Count\nAlice,10,2\nBob,20,3', 'a.csv', header=[1, 2])
    assert list(source.frame.columns) == ['Customer Name', 'Sales Amount', 'Sales Count']
    assert source.header == 2 and source.header_end == 3
    assert list(source.frame.index) == [4, 5]
    assert source.frame.iloc[0].tolist() == ['Alice', '10', '2']


def test_blank_and_repeated_header_parts():
    source = load(b'ID,,Order\nID,Name, Date \n1,Alice,2026-01-01', 'a.csv', header=[0, 1])
    assert list(source.frame.columns) == ['ID', 'Name', 'Order Date']
    source = load(b'" Customer\n Name ",Total\nAlice,10', 'a.csv')
    assert list(source.frame.columns) == ['Customer Name', 'Total']


@pytest.mark.parametrize('header', [[1, 0], [0, 2], [], [0, 10], [-1, 0], '1-2'])
def test_invalid_header_ranges(header):
    with pytest.raises(ValueError):
        load(b'A,B\nC,D\n1,2', 'a.csv', header=header)


@pytest.mark.parametrize('content', [b'A,A\nB,B\n1,2', b'A,\nB,\n1,2'])
def test_duplicate_or_blank_after_combining_blocks(content):
    with pytest.raises(ValueError):
        load(content, 'a.csv', header=[0, 1])


def merged_workbook():
    book = Workbook()
    ws = book.active
    ws.title = 'Sales'
    ws.append(['Customer', 'Sales', None])
    ws.append([None, 'Amount', 'Count'])
    ws.merge_cells('A1:A2')
    ws.merge_cells('B1:C1')
    ws.append(['Alice', 10, 2])
    ws.append(['Bob', 'bad', 3])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_excel_actual_merged_cells_and_template_reuse():
    content = merged_workbook()
    source = load(content, 'sales.xlsx', 'Sales', [0, 1])
    assert list(source.frame.columns) == ['Customer', 'Sales Amount', 'Sales Count']
    config = default_config()
    config.update(trace=True, columns=[{'name': 'Customer', 'type': 'Text', 'required': True},
                                      {'name': 'Sales Amount', 'type': 'Number', 'required': False},
                                      {'name': 'Sales Count', 'type': 'Number', 'required': False}])
    config['headers'] = {'sales.xlsx': {'Sales': [0, 1]}}
    config['mappings'] = {source.key: {h: h for h in source.frame.columns}}
    restored = read_config(json.dumps(config))
    reloaded = load(content, 'sales.xlsx', 'Sales', restored['headers']['sales.xlsx']['Sales'])
    master, issues, summary = consolidate([reloaded], restored, restored['mappings'])
    assert not summary['blocked'] and len(master) == 2
    assert master['Source Row'].tolist() == [3, 4]
    assert issues.iloc[0]['Source row'] == 4
    assert issues.iloc[0]['Source column'] == 'Sales Amount'
